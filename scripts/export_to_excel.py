#!/usr/bin/env python
"""
Export enrichment data to Excel.

Reads from local SQLite (--db) or Supabase (default) and writes one tab per
table into a single .xlsx file.

Usage:
    # From local SQLite (test mode)
    python scripts/export_to_excel.py --db local_test.db --out output.xlsx

    # From Supabase
    python scripts/export_to_excel.py --out output.xlsx --sector Retailers --country AE

    # From Supabase, specific sector/country
    python scripts/export_to_excel.py --out retailers_AE.xlsx --sector Retailers --country AE
"""
from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, "src")

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_ALT_FILL = PatternFill("solid", fgColor="EEF2F7")


# ---------------------------------------------------------------------------
# Data fetchers
# ---------------------------------------------------------------------------


def _fetch_sqlite(db_path: str) -> dict[str, list[dict[str, Any]]]:
    """Load all four tables from a local SQLite file."""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    _json_cols = {
        "sub_sector_tags", "sector_metadata", "raw_search_results",
        "raw_website_data", "raw_llm_extraction", "dimension_scores",
        "confidence_bands", "executives", "request_data", "response_data",
    }

    def _rows(table: str) -> list[dict[str, Any]]:
        try:
            cur = conn.execute(f"SELECT * FROM {table}")  # noqa: S608
        except sqlite3.OperationalError:
            return []
        cols = [d[0] for d in cur.description]
        result = []
        for row in cur.fetchall():
            d: dict[str, Any] = {}
            for col, val in zip(cols, row):
                if col in _json_cols and isinstance(val, str):
                    try:
                        val = json.loads(val)
                    except (json.JSONDecodeError, TypeError):
                        pass
                d[col] = val
            result.append(d)
        return result

    data = {
        "companies": _rows("companies"),
        "company_details": _rows("company_details"),
        "company_scores": _rows("company_scores"),
        "enrichment_audit": _rows("enrichment_audit"),
    }
    conn.close()
    return data


def _fetch_supabase(sector: str, country: str) -> dict[str, list[dict[str, Any]]]:
    """Load data from Supabase, filtered by sector/country."""
    from hak_talent_mapping.config import Settings
    from supabase import create_client

    s = Settings()  # type: ignore[call-arg]
    client = create_client(s.supabase_url, s.supabase_key)

    def _paginate(table: str, filters: dict[str, str]) -> list[dict[str, Any]]:
        results: list[dict[str, Any]] = []
        offset = 0
        page = 1000
        while True:
            q = client.table(table).select("*")
            for k, v in filters.items():
                q = q.eq(k, v)
            resp = q.range(offset, offset + page - 1).execute()
            if not resp.data:
                break
            results.extend(resp.data)
            if len(resp.data) < page:
                break
            offset += page
        return results

    companies = _paginate("companies", {"sector": sector})
    details = _paginate("company_details", {"sector": sector, "country_code": country})
    detail_ids = [r["id"] for r in details]

    scores: list[dict[str, Any]] = []
    audit: list[dict[str, Any]] = []
    # Fetch scores + audit in batches of 100 IDs (Supabase in() limit)
    for i in range(0, len(detail_ids), 100):
        batch = detail_ids[i : i + 100]
        s_resp = client.table("company_scores").select("*").in_("company_detail_id", batch).execute()
        scores.extend(s_resp.data or [])
        a_resp = client.table("enrichment_audit").select("*").in_("company_detail_id", batch).execute()
        audit.extend(a_resp.data or [])

    return {
        "companies": companies,
        "company_details": details,
        "company_scores": scores,
        "enrichment_audit": audit,
    }


# ---------------------------------------------------------------------------
# Excel writer
# ---------------------------------------------------------------------------


def _flatten_value(val: Any) -> str | int | float | None:
    """Flatten lists/dicts to compact JSON strings for cell display."""
    if val is None:
        return None
    if isinstance(val, (int, float, bool)):
        return val
    if isinstance(val, (list, dict)):
        return json.dumps(val, ensure_ascii=False)
    return str(val)


def _write_sheet(
    ws: Any,
    rows: list[dict[str, Any]],
    skip_cols: set[str] | None = None,
) -> None:
    """Write rows to a worksheet with styled headers and alternating row colours."""
    if not rows:
        ws.append(["(no data)"])
        return

    skip = skip_cols or set()
    headers = [k for k in rows[0].keys() if k not in skip]

    # Header row
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=False)

    # Data rows
    for row_idx, record in enumerate(rows, 2):
        fill = _ALT_FILL if row_idx % 2 == 0 else None
        for col_idx, header in enumerate(headers, 1):
            raw = record.get(header)
            val = _flatten_value(raw)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(wrap_text=False)

    # Auto-width (cap at 60)
    for col_idx, header in enumerate(headers, 1):
        max_len = len(header)
        for row_idx in range(2, min(len(rows) + 2, 52)):  # sample first 50 rows
            v = ws.cell(row=row_idx, column=col_idx).value
            if v:
                max_len = max(max_len, min(len(str(v)), 60))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 2

    ws.freeze_panes = "A2"


_HEAVY_COLS = {"raw_search_results", "raw_website_data", "raw_llm_extraction",
               "request_data", "response_data"}


def _write_leadership_sheet(wb: Any, details: list[dict[str, Any]]) -> None:
    """Dedicated tab: one row per leader extracted from sector_metadata."""
    ws = wb.create_sheet("leadership")
    headers = ["company_name", "company_id", "sector", "leader_name", "leader_title"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    row_idx = 2
    for record in details:
        sm = record.get("sector_metadata") or {}
        if isinstance(sm, str):
            try:
                sm = json.loads(sm)
            except Exception:
                sm = {}
        leaders = sm.get("leadership_names", [])
        if not leaders:
            continue
        for leader in leaders:
            ws.cell(row=row_idx, column=1, value=record.get("name"))
            ws.cell(row=row_idx, column=2, value=record.get("company_id"))
            ws.cell(row=row_idx, column=3, value=record.get("sector"))
            ws.cell(row=row_idx, column=4, value=leader.get("name") if isinstance(leader, dict) else str(leader))
            ws.cell(row=row_idx, column=5, value=leader.get("title") if isinstance(leader, dict) else "")
            row_idx += 1

    for col_idx in range(1, 6):
        ws.column_dimensions[get_column_letter(col_idx)].width = 30
    ws.freeze_panes = "A2"


def export(
    data: dict[str, list[dict[str, Any]]],
    out_path: str,
) -> None:
    wb = openpyxl.Workbook()

    # Tab 1: companies (listing data)
    ws_companies = wb.active
    ws_companies.title = "companies"
    _write_sheet(ws_companies, data["companies"])

    # Tab 2: company_details (enrichment profiles, skip heavy raw blobs)
    ws_details = wb.create_sheet("company_details")
    _write_sheet(ws_details, data["company_details"], skip_cols=_HEAVY_COLS)

    # Tab 3: leadership (exploded from sector_metadata)
    _write_leadership_sheet(wb, data["company_details"])

    # Tab 4: company_scores
    ws_scores = wb.create_sheet("company_scores")
    _write_sheet(ws_scores, data["company_scores"])

    # Tab 5: enrichment_audit (skip heavy blobs)
    ws_audit = wb.create_sheet("enrichment_audit")
    _write_sheet(ws_audit, data["enrichment_audit"], skip_cols=_HEAVY_COLS)

    wb.save(out_path)
    print(f"Saved: {out_path}")
    for name, rows in data.items():
        print(f"  {name}: {len(rows)} rows")
    leadership_count = sum(
        len((r.get("sector_metadata") or {}).get("leadership_names", []))
        if isinstance(r.get("sector_metadata"), dict)
        else 0
        for r in data["company_details"]
    )
    print(f"  leadership: {leadership_count} leaders")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(description="Export enrichment data to Excel")
    parser.add_argument(
        "--db",
        type=str,
        default=None,
        metavar="PATH",
        help="Read from local SQLite file (test mode). If omitted, reads from Supabase.",
    )
    parser.add_argument(
        "--out",
        type=str,
        default=None,
        metavar="PATH",
        help="Output .xlsx path (default: export_<timestamp>.xlsx)",
    )
    parser.add_argument(
        "--sector",
        type=str,
        default="Retailers",
        help="(Supabase only) Filter by sector (default: Retailers)",
    )
    parser.add_argument(
        "--country",
        type=str,
        default="AE",
        help="(Supabase only) Filter by country code (default: AE)",
    )
    args = parser.parse_args()

    out = args.out or f"export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    if args.db:
        db_path = str(Path(args.db).resolve())
        if not Path(db_path).exists():
            print(f"Error: SQLite file not found: {db_path}", file=sys.stderr)
            sys.exit(1)
        print(f"Reading from local SQLite: {db_path}")
        data = _fetch_sqlite(db_path)
    else:
        print(f"Reading from Supabase (sector={args.sector}, country={args.country})…")
        data = _fetch_supabase(sector=args.sector, country=args.country)

    export(data, out)


if __name__ == "__main__":
    main()
